"""Render KPI scorecard + Markdown narrative into a client-ready .docx (US Letter)."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .narrate import month_name

ACCENT = RGBColor(0x1F, 0x4E, 0x5A)
SCORECARD = [
    ("Spend", "spend"), ("Conversions", "conversions"), ("CPA", "cpa"),
    ("Revenue", "revenue"), ("ROAS", "roas"), ("CTR", "ctr"), ("Conversion rate", "cvr"),
]


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _add_runs(paragraph, text: str) -> None:
    """Supports **bold** spans."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            paragraph.add_run(part).bold = i % 2 == 1


def render(kpis: dict, narrative_md: str, out_path: str | Path) -> Path:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading(f"{kpis['client']}: Paid Media Performance", level=0)
    title.runs[0].font.color.rgb = ACCENT
    doc.add_paragraph(f"{month_name(kpis['month'])} vs. {month_name(kpis['prior_month'])}").runs[0].italic = True

    t = kpis["totals"]["display"]
    rows = [(label, key) for label, key in SCORECARD if t["current"][key] is not None]
    table = doc.add_table(rows=1 + len(rows), cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["Metric", month_name(kpis["month"]), month_name(kpis["prior_month"]), "MoM"]):
        cell = table.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(cell, "1F4E5A")
    for r, (label, key) in enumerate(rows, start=1):
        vals = [label, t["current"][key], t["prior"][key], t["change"][key] or "n/a"]
        for c, v in enumerate(vals):
            table.rows[r].cells[c].text = v
    doc.add_paragraph()

    for line in narrative_md.splitlines():
        line = line.rstrip()
        if not line:
            continue
        if line.startswith("## "):
            h = doc.add_heading(line[3:], level=1)
            h.runs[0].font.color.rgb = ACCENT
        elif line.startswith(("- ", "* ")):
            _add_runs(doc.add_paragraph(style="List Bullet"), line[2:])
        else:
            _add_runs(doc.add_paragraph(), line)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path
